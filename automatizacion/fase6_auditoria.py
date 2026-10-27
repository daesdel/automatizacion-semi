#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
FASE 6 - Auditoría de tickets huérfanos.

Busca IRs presentes en DATA IN (JIRA), drive y seguimiento.xlsx que NO estén
en la columna A de PREVISIONES JUAN, calcula sus columnas H, AB y AJ y genera
Auditoria_Tickets_Faltantes.xlsx.

Uso:
    python automatizacion/fase6_auditoria.py
"""

import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl
import xlsxwriter

HERE = Path(__file__).resolve().parent
BASE_DIR = HERE.parent / "Nueva carpeta (2)"


def _arg_path(nombre, defecto):
    pref = f"--{nombre}="
    for a in sys.argv[1:]:
        if a.startswith(pref):
            return Path(a.split("=", 1)[1]).expanduser()
    return defecto


JIRA = _arg_path("jira", BASE_DIR / "JIRA.xlsx")
SEGUIMIENTO = _arg_path("seguimiento", BASE_DIR / "resultados" / "seguimiento.xlsx")
HOJA = _arg_path("hoja", BASE_DIR / "Hoja de cálculo sin título.xlsx")
PREV = _arg_path("prev", BASE_DIR / "resultados" / "Seguimiento_Previsiones_Limpio.xlsx")
OUT = _arg_path("out", BASE_DIR / "resultados" / "Auditoria_Tickets_Faltantes.xlsx")

sys.path.insert(0, str(HERE))
import fase5_previsiones as f5  # noqa: E402


def _claves_jira():
    """Devuelve (IR de radio, IR de radio con Fecha PO > 0)."""
    wb = openpyxl.load_workbook(JIRA, read_only=True, data_only=True)
    ws = wb["IR_VDF"]
    hdr = list(next(ws.iter_rows(min_row=4, max_row=4, values_only=True)))
    idx = {h: i for i, h in enumerate(hdr) if h}
    i_clave, i_po = idx.get("Clave"), idx.get("Fecha PO")

    radio, con_po = set(), set()
    for fila in ws.iter_rows(min_row=5, values_only=True):
        clave = fila[i_clave] if i_clave is not None and i_clave < len(fila) else None
        if clave in (None, ""):
            continue
        clave = str(clave).strip()
        if not clave.upper().startswith("IR-"):
            continue
        radio.add(clave)
        po = fila[i_po] if i_po is not None and i_po < len(fila) else None
        if isinstance(po, (date, datetime)):
            con_po.add(clave)
        elif isinstance(po, (int, float)) and po > 0:
            con_po.add(clave)
    wb.close()
    return radio, con_po


def _claves_seguimiento():
    wb = openpyxl.load_workbook(SEGUIMIENTO, read_only=True, data_only=True)
    ws = wb["NUEVO SEGUIMIENTO"]
    out = set()
    for fila in ws.iter_rows(min_row=2, min_col=2, max_col=2, values_only=True):
        if fila[0] not in (None, ""):
            ir = str(fila[0]).strip()
            if ir.upper().startswith("IR-"):
                out.add(ir)
    wb.close()
    return out


def _claves_prev():
    if not PREV.is_file():
        return set()
    wb = openpyxl.load_workbook(PREV, read_only=True, data_only=True)
    ws = wb["PREVISIONES JUAN"]
    out = set()
    for fila in ws.iter_rows(min_row=2, min_col=1, max_col=1, values_only=True):
        if fila[0] not in (None, ""):
            out.add(str(fila[0]).strip())
    wb.close()
    return out


def _validar_jira() -> None:
    if not JIRA.is_file():
        sys.exit(
            "No existe JIRA.xlsx en la ruta indicada. Ejecuta primero el "
            "proceso principal (Fase 1) para generarlo.\n"
            f"Ruta buscada: {JIRA}"
        )
    wb = openpyxl.load_workbook(JIRA, read_only=True, data_only=True)
    hojas = wb.sheetnames
    wb.close()
    if "IR_VDF" not in hojas:
        sys.exit(
            "El archivo indicado no parece un JIRA.xlsx válido: no contiene "
            "la pestaña IR_VDF. Ejecuta primero el proceso principal.\n"
            f"Archivo: {JIRA}\nHojas: {', '.join(hojas)}"
        )


def main() -> None:
    if not JIRA.is_file():
        sys.exit(
            "No existe JIRA.xlsx en la ruta indicada. Ejecuta primero el "
            "proceso principal (Fase 1) para generarlo.\n"
            f"Ruta buscada: {JIRA}"
        )
    _validar_jira()
    faltan = [p for p in (SEGUIMIENTO, HOJA) if not p.is_file()]
    if faltan:
        sys.exit("Faltan archivos: " + ", ".join(str(p) for p in faltan))

    print("=" * 70)
    print("  FASE 6 · Auditoría de tickets faltantes")
    print("=" * 70)

    prev = _claves_prev()
    jira_radio, jira_po = _claves_jira()
    drive = _claves_seguimiento()

    universo = jira_radio | drive
    elegibles = {ir for ir in universo if ir in drive or ir in jira_po}
    origenes = {}
    for ir in elegibles:
        if ir in drive:
            origenes.setdefault(ir, set()).add("drive")
        if ir in jira_po:
            origenes.setdefault(ir, set()).add("DATA IN (Fecha PO>0)")

    huerfanos = sorted(ir for ir in elegibles if ir not in prev)
    print(f"  PREVISIONES: {len(prev):,} IRs")
    print(f"  IR de radio en DATA IN: {len(jira_radio):,} | con Fecha PO>0: {len(jira_po):,}")
    print(f"  IR de radio en drive: {len(drive):,} | elegibles: {len(elegibles):,}")
    print(f"  Huérfanos: {len(huerfanos):,}")

    jmap = f5.leer_jira_map()
    tx, act = f5.leer_subtareas()

    wb = xlsxwriter.Workbook(str(OUT))
    ws = wb.add_worksheet("FALTANTES")
    f_cab = wb.add_format({"bold": True, "bg_color": "#D9E1F2", "border": 1})
    f_txt = wb.add_format({"border": 1, "valign": "top"})
    cab = ["IR", "Origen", "H", "AB", "AJ"]
    for c, t in enumerate(cab):
        ws.write_string(0, c, t, f_cab)
    ws.set_column(0, 0, 14)
    ws.set_column(1, 1, 24)
    ws.set_column(2, 2, 42)
    ws.set_column(3, 4, 16)
    for r, ir in enumerate(huerfanos, start=1):
        h, ab, aj = f5.calcular_manual(ir, jmap, tx, act)
        ws.write_string(r, 0, ir, f_txt)
        ws.write_string(r, 1, " + ".join(sorted(origenes[ir])), f_txt)
        ws.write_string(r, 2, h or "", f_txt)
        ws.write_string(r, 3, ab or "", f_txt)
        ws.write_string(r, 4, aj or "", f_txt)
    wb.close()
    print(f"  [OK] {OUT} ({OUT.stat().st_size / 1e6:.2f} MB)")


if __name__ == "__main__":
    main()