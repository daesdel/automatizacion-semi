#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Genera plantilla_previsiones.xlsx (una sola vez).

Contiene:
  - PREVISIONES JUAN: constantes, fórmulas y formato exacto (87 columnas, sin CH).
  - Solo PREVISIONES JUAN (sin DATA IN).

A partir de aquí, fase5_previsiones.py ya no necesita el libro original.
"""

import sys
import time
from pathlib import Path
import re as _re
import xlsxwriter.worksheet as _xw


class FormulasNormales:
    """Evita que xlsxwriter convierta XLOOKUP/SWITCH en fórmulas matriciales."""

    def __enter__(self):
        self._orig = _xw.re_dynamic_function
        _xw.re_dynamic_function = _re.compile(r"(?!x)x")
        return self

    def __exit__(self, *args):
        _xw.re_dynamic_function = self._orig


import openpyxl
import xlsxwriter
from openpyxl.worksheet.formula import ArrayFormula

HERE = Path(__file__).resolve().parent
BASE_DIR = HERE.parent / "Nueva carpeta (2)"
ORIGEN = BASE_DIR / "SEGUIMIENTO PREVISIONES.xlsx"
DESTINO = BASE_DIR / "plantilla_previsiones.xlsx"

sys.path.insert(0, str(HERE))
import estilos_previsiones as est  # noqa: E402

COL_CH = 86


def escribir(ws, fila, col, valor, fmt, fecha_fmt):
    if valor is None or valor == "":
        return
    if isinstance(valor, ArrayFormula):
        ws.write_array_formula(fila, col, fila, col, valor.text, fmt)
    elif isinstance(valor, str) and valor.startswith("="):
        ws.write_formula(fila, col, valor, fmt)
    else:
        from datetime import date, datetime
        if isinstance(valor, datetime):
            ws.write_datetime(fila, col, valor, fmt or fecha_fmt)
        elif isinstance(valor, date):
            ws.write_datetime(fila, col, datetime(valor.year, valor.month, valor.day),
                              fmt or fecha_fmt)
        elif isinstance(valor, (int, float)):
            ws.write_number(fila, col, valor, fmt)
        else:
            ws.write_string(fila, col, str(valor), fmt)


def main() -> None:
    t0 = time.time()
    if not ORIGEN.is_file():
        sys.exit(f"No existe {ORIGEN}")

    wb_src = openpyxl.load_workbook(ORIGEN, read_only=True, data_only=False)
    wb = xlsxwriter.Workbook(str(DESTINO), {"constant_memory": True})
    fmt_fecha = wb.add_format({"num_format": "dd/mm/yyyy"})
    cab, dat = est.crear_formatos(wb)

    # ---------------- PREVISIONES JUAN ----------------
    ws = wb.add_worksheet("PREVISIONES JUAN")
    for col, ancho in est.ANCHOS.items():
        ws.set_column(col - 1, col - 1, ancho)
    ws.set_row(0, est.ALTO_FILA1)
    ws.freeze_panes(1, 1)

    filas = 0
    with FormulasNormales():
        for r, fila in enumerate(wb_src["PREVISIONES JUAN"].iter_rows(values_only=True)):
            destino = 0
            for idx, valor in enumerate(fila, start=1):
                if idx == COL_CH:
                    continue
                destino += 1
                fmt = cab[destino] if r == 0 else dat[destino]
                manual = destino in (1, 8, 28, 36)  # A, H, AB, AJ
                if (r > 0 and not manual
                        and not isinstance(valor, ArrayFormula)
                        and not (isinstance(valor, str) and valor.startswith("="))
                        and valor in (None, "")):
                    valor = "AUTO"
                if valor == "AUTO":
                    ws.write_string(r, destino - 1, "AUTO", fmt)
                else:
                    escribir(ws, r, destino - 1, valor, fmt, fmt_fecha)
            filas = r + 1

    # Formato condicional (7 estados)
    verde = wb.add_format({"bg_color": "#92D050"})
    rojo = wb.add_format({"bg_color": "#FF0000"})
    gris = wb.add_format({"bg_color": "#8D8B8B"})
    reglas = [("Aceptado Zona", verde), ("RIP", gris), ("Doc & Inv Radio", verde),
              ("KPIS OK", verde), ("ON AIR", verde), ("RFOA", verde),
              ("Parada de Reloj", rojo)]
    for rango, ref in [(f"A2:BH{filas}", 2), ("AJ1314", 1314), ("AJ3374:AJ3376", 3374)]:
        for texto, fmt in reglas:
            ws.conditional_format(rango, {
                "type": "formula",
                "criteria": f'=$N{ref}="{texto}"',
                "format": fmt,
            })

    wb.close()
    wb_src.close()
    print(f"[OK] {DESTINO.name} ({DESTINO.stat().st_size/1e6:.1f} MB) "
          f"con {filas:,} filas en PREVISIONES JUAN, {time.time()-t0:.1f}s")


if __name__ == "__main__":
    main()