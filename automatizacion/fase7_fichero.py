#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
FASE 7 - Fichero de trabajo (equivalente a copiar_juan).

Genera [DDMMYYYY]_FICHERO.xlsx con:
  - PREVISIONES JUAN (valores y formato, sin fórmulas)
  - RFOA SIN PREVISION (tabla Tabla1 de Previsiones de RFOA_central.xlsx,
    solo celdas visibles)
  - DUPLICADOS (IRs repetidos de la columna A, pintados en morado)

Uso:
    python automatizacion/fase7_fichero.py
    python automatizacion/fase7_fichero.py --prev=RUTA --rfoa=RUTA --out=RUTA
"""

import shutil
import sys
import time
import zipfile
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import coordinate_from_string
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = Path(__file__).resolve().parent
BASE_DIR = HERE.parent / "Nueva carpeta (2)"
RESULTADOS = BASE_DIR / "resultados"


def _arg_path(nombre, defecto):
    pref = f"--{nombre}="
    for a in sys.argv[1:]:
        if a.startswith(pref):
            return Path(a.split("=", 1)[1]).expanduser()
    return defecto


PREV = _arg_path("prev", RESULTADOS / "Seguimiento_Previsiones_Limpio.xlsx")
RFOA = _arg_path("rfoa", BASE_DIR / "Previsiones de RFOA_central.xlsx")
OUT = _arg_path("out", RESULTADOS / f"{date.today():%d%m%Y}_FICHERO.xlsx")


def _comprobar_xlsx(path: Path, etiqueta: str) -> None:
    """Verifica que el .xlsx es un ZIP válido; reintenta si está escribiéndose."""
    for intento in range(5):
        try:
            with zipfile.ZipFile(path) as z:
                malo = z.testzip()
            if malo:
                raise zipfile.BadZipFile(f"entrada dañada: {malo}")
            return
        except (zipfile.BadZipFile, EOFError, OSError):
            if intento == 4:
                sys.exit(
                    f"El archivo {etiqueta} está dañado o abierto en Excel. "
                    f"Cierra Excel y vuelve a intentarlo.\nRuta: {path}"
                )
            time.sleep(1)


MORADO = PatternFill("solid", fgColor="800080")
BLANCO = Font(color="FFFFFF", bold=True)


def _ref_limites(ref):
    ini, fin = ref.split(":")
    ws_ini = coordinate_from_string(ini)
    ws_fin = coordinate_from_string(fin)
    return (column_index_from_string(ws_ini[0]), ws_ini[1],
            column_index_from_string(ws_fin[0]), ws_fin[1])


def copiar_previsiones(wb, wv):
    ws = wb["PREVISIONES JUAN"]
    wsv = wv["PREVISIONES JUAN"]
    for fila in ws.iter_rows():
        for celda in fila:
            v = celda.value
            if isinstance(v, ArrayFormula) or (isinstance(v, str) and v.startswith("=")):
                celda.value = wsv[celda.coordinate].value
    for nombre in list(wb.sheetnames):
        if nombre != "PREVISIONES JUAN":
            wb.remove(wb[nombre])
    return ws


def copiar_rfoa(wb, ws_prev):
    if not RFOA.is_file():
        print(f"    [AVISO] No existe {RFOA.name}; se omite RFOA SIN PREVISION.")
        return
    try:
        origen = openpyxl.load_workbook(RFOA, data_only=True)
    except (zipfile.BadZipFile, EOFError, OSError):
        sys.exit(
            "El archivo RFOA central está dañado o abierto en Excel. "
            f"Cierra Excel y vuelve a intentarlo.\nRuta: {RFOA}"
        )
    if "Nodos" not in origen.sheetnames:
        print("    [AVISO] El RFOA central no tiene hoja Nodos; se omite.")
        origen.close()
        return
    ws_o = origen["Nodos"]
    tabla = ws_o.tables.get("Tabla1")
    if tabla is None:
        print("    [AVISO] No existe Tabla1 en Nodos; se omite.")
        origen.close()
        return

    min_col, min_row, max_col, max_row = _ref_limites(tabla.ref)
    filas = []
    for r in range(min_row, max_row + 1):
        if r != min_row and ws_o.row_dimensions[r].hidden:
            continue
        filas.append([ws_o.cell(row=r, column=c).value
                      for c in range(min_col, max_col + 1)])

    ws = wb.create_sheet("RFOA SIN PREVISION")
    for r, fila in enumerate(filas, start=1):
        for c, v in enumerate(fila, start=1):
            ws.cell(row=r, column=c, value=v)
    for c in range(1, max_col - min_col + 2):
        ws.column_dimensions[get_column_letter(c)].width = 18
    for c in range(1, len(filas[0]) + 1):
        celda = ws.cell(row=1, column=c)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1F4E78")
        celda.alignment = Alignment(horizontal="center", vertical="center")

    if filas:
        ref = f"A1:{get_column_letter(len(filas[0]))}{len(filas)}"
        try:
            t = Table(displayName="tbl_RFOA_SIN_PREVISION", ref=ref)
            t.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(t)
        except Exception as e:  # noqa: BLE001
            print(f"    [AVISO] No se pudo crear la tabla RFOA: {e}")
    print(f"    RFOA SIN PREVISION: {len(filas) - 1:,} filas visibles")
    origen.close()


def generar_duplicados(wb, ws_prev):
    ws_d = wb.create_sheet("DUPLICADOS")
    ws_d["A1"] = "TICKET"
    ws_d["B1"] = "HOJA"
    ws_d["A1"].font = Font(bold=True)
    ws_d["B1"].font = Font(bold=True)
    ws_d.column_dimensions["A"].width = 18
    ws_d.column_dimensions["B"].width = 22

    uf = ws_prev.max_row
    vistos = {}
    for r in range(2, uf + 1):
        v = ws_prev.cell(row=r, column=1).value
        if v not in (None, ""):
            vistos[str(v).strip()] = vistos.get(str(v).strip(), 0) + 1

    dup = {k for k, n in vistos.items() if n > 1}
    f = 2
    for r in range(2, uf + 1):
        v = ws_prev.cell(row=r, column=1).value
        if v is not None and str(v).strip() in dup:
            ws_prev.cell(row=r, column=1).fill = MORADO
            ws_prev.cell(row=r, column=1).font = BLANCO
            ws_d.cell(row=f, column=1, value=v)
            ws_d.cell(row=f, column=2, value="PREVISIONES JUAN")
            f += 1
    print(f"    DUPLICADOS: {len(dup):,} IRs repetidos ({f - 2:,} filas)")


def main() -> None:
    print("=" * 70)
    print("  FASE 7 · Generando FICHERO de trabajo")
    print("=" * 70)
    if not PREV.is_file():
        sys.exit(f"No existe el PREVISIONES de origen: {PREV}")
    if not RFOA.is_file():
        sys.exit(f"No existe el RFOA central: {RFOA}")
    _comprobar_xlsx(PREV, "PREVISIONES de origen")
    _comprobar_xlsx(RFOA, "RFOA central")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(PREV, OUT)
    print(f"    Base: {PREV.name} -> {OUT.name}")

    wb = openpyxl.load_workbook(OUT)
    wv = openpyxl.load_workbook(PREV, data_only=True)

    ws_prev = copiar_previsiones(wb, wv)
    copiar_rfoa(wb, ws_prev)
    generar_duplicados(wb, ws_prev)

    wb.save(OUT)
    wv.close()
    wb.close()
    print(f"\n  [OK] {OUT} ({OUT.stat().st_size / 1e6:.2f} MB)")


if __name__ == "__main__":
    main()